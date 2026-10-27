import os
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import pymupdf  # PyMuPDF
import docx

from app.config import config
from app.database import get_db, DatabaseService
from app.database.models import DocumentRecord
from app.core.ai_manager import ai_manager


class DocumentService:
    """Service for processing PDFs, DOCX, and text files with AI intelligence."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".log"}

    def __init__(self, db: Optional[DatabaseService] = None):
        self.db = db or get_db()

    def is_supported(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in self.SUPPORTED_EXTENSIONS

    def extract_text(self, file_path: str) -> Tuple[str, str]:
        """Extract plain text from local document.
        
        Returns:
            Tuple of (extracted_text, file_type)
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}. Supported: {', '.join(self.SUPPORTED_EXTENSIONS)}")

        extracted_text = ""

        if ext == ".pdf":
            try:
                doc = pymupdf.open(str(path))
                if doc.is_encrypted:
                    raise ValueError("Cannot extract text: The PDF file is encrypted/password protected.")
                pages_text = []
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    pages_text.append(page.get_text("text"))
                extracted_text = "\n\n".join(pages_text).strip()
                doc.close()
                if not extracted_text:
                    raise ValueError(
                        "No readable text could be extracted from this PDF. "
                        "The document may consist solely of scanned images or raster graphics."
                    )
            except Exception as e:
                if isinstance(e, ValueError):
                    raise
                raise RuntimeError(f"Error parsing PDF file: {str(e)}")

        elif ext == ".docx":
            try:
                doc = docx.Document(str(path))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also extract text from tables
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            paragraphs.append(" | ".join(row_text))
                extracted_text = "\n\n".join(paragraphs).strip()
                if not extracted_text:
                    raise ValueError("The Word document (.docx) does not contain extractable text.")
            except Exception as e:
                if isinstance(e, ValueError):
                    raise
                raise RuntimeError(f"Error reading DOCX file: {str(e)}")

        else:  # .txt, .md, .log
            try:
                with open(str(path), "r", encoding="utf-8", errors="replace") as f:
                    extracted_text = f.read().strip()
                if not extracted_text:
                    raise ValueError("The text file is empty.")
            except Exception as e:
                raise RuntimeError(f"Error reading text file: {str(e)}")

        return extracted_text, ext.replace(".", "")

    def ingest_document(self, file_path: str) -> DocumentRecord:
        """Extract text and persist metadata into SQLite repository."""
        path = Path(file_path)
        extracted_text, file_type = self.extract_text(str(path))
        file_size = path.stat().st_size

        record = self.db.add_document(
            filename=path.name,
            file_path=str(path.resolve()),
            file_type=file_type,
            file_size=file_size,
            extracted_text=extracted_text,
            summary=None
        )
        return record

    def summarize_document(self, doc_record: DocumentRecord) -> str:
        """Generate executive summary and key points using configured AI."""
        # Truncate text context safely if very long to stay within standard context limits
        context = doc_record.extracted_text[:12000]
        prompt = (
            f"Please analyze the following document titled '{doc_record.filename}'.\n"
            "Provide:\n"
            "1. An executive summary (2-3 paragraphs).\n"
            "2. 5-7 key takeaways or actionable points in bullet format.\n"
            "3. Any notable dates, figures, or decisions.\n\n"
            f"--- DOCUMENT CONTENT ---\n{context}\n--- END CONTENT ---"
        )
        messages = [{"role": "user", "content": prompt}]
        summary, _ = ai_manager.generate(messages)
        
        # Cache summary in database
        with self.db.get_connection() as conn:
            conn.execute("UPDATE documents SET summary = ? WHERE id = ?", (summary, doc_record.id))
            conn.commit()
        doc_record.summary = summary
        return summary

    def answer_question(self, doc_record: DocumentRecord, question: str) -> str:
        """Answer user query grounded in document context."""
        context = doc_record.extracted_text[:12000]
        messages = [{"role": "user", "content": question}]
        response, _ = ai_manager.generate(messages, document_context=context)
        return response

    def list_documents(self) -> List[DocumentRecord]:
        return self.db.list_documents()

    def get_document(self, doc_id: str) -> Optional[DocumentRecord]:
        return self.db.get_document(doc_id)

    def delete_document(self, doc_id: str):
        self.db.delete_document(doc_id)


document_service = DocumentService()
